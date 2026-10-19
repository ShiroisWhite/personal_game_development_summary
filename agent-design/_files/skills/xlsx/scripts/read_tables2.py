import sys
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

def read_table(filepath, max_rows=20):
    print(f"\n{'='*60}")
    print(f"=== {filepath.split(chr(92))[-1]} ===")
    wb = load_workbook(filepath, data_only=True)
    print(f"Sheets: {wb.sheetnames}")
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n--- Sheet: {name} (rows={ws.max_row}, cols={ws.max_column}) ---")
        for i, row in enumerate(ws.iter_rows(max_row=min(max_rows, ws.max_row), values_only=False), 1):
            vals = [str(c.value) if c.value is not None else "" for c in row]
            print(f"  Row{i}: {vals}")
    wb.close()

if __name__ == "__main__":
    files = [
        "<config_root>/BuildInfo.xlsx",
        "<config_root>/Mews.xlsx",
        "<config_root>/AnimalTame.xlsx",
        "<config_root>/Formula.xlsx",
    ]
    for f in files:
        try:
            read_table(f)
        except Exception as e:
            print(f"ERROR reading {f}: {e}")
