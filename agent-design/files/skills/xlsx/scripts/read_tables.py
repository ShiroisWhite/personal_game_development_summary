import sys
from openpyxl import load_workbook

def read_table(filepath, max_rows=25):
    print(f"\n{'='*60}")
    print(f"=== {filepath.split(chr(92))[-1]} ===")
    wb = load_workbook(filepath, data_only=True)
    print(f"Sheets: {wb.sheetnames}")
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n--- Sheet: {name} (rows={ws.max_row}, cols={ws.max_column}) ---")
        for i, row in enumerate(ws.iter_rows(max_row=min(max_rows, ws.max_row), values_only=False), 1):
            vals = [str(c.value) if c.value is not None else "" for c in row]
            print(f"  Row{i}: {vals}")
    wb.close()

if __name__ == "__main__":
    files = [
        "<config_root>/WorkBenchUnit.xlsx",
        "<config_root>/Planting.xlsx",
        "<config_root>/Husbandry.xlsx",
        "<config_root>/HusbandryHouse.xlsx",
    ]
    for f in files:
        try:
            read_table(f)
        except Exception as e:
            print(f"ERROR reading {f}: {e}")
