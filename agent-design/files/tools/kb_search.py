#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""在当前内容底座中做全文检索。

用法：
    python tools/kb_search.py 关键词 [关键词 ...]
    python tools/kb_search.py 关键词 --base soc-general
    python tools/kb_search.py 关键词 --ext .md --limit 50 --context 1
    python tools/kb_search.py 关键词 --no-excel        # 不扫 Excel 单元格
    python tools/kb_search.py --files

默认：所有关键词都必须命中（AND）。加 --any 改为任一命中。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

TOOLS_DIR = Path(__file__).resolve().parent
ROOT = TOOLS_DIR.parent
REGISTRY_PATH = ROOT / "bases" / "registry.json"

DEFAULT_EXTS = {
    ".md", ".txt", ".json", ".jsonc", ".csv", ".tsv",
    ".yaml", ".yml", ".py", ".js", ".ts", ".html", ".css",
    ".sql", ".ini", ".toml", ".xml",
}
SKIP_DIRS = {".git", ".svn", "node_modules", "__pycache__", ".venv", "venv", "dist", "build"}
EXCEL_EXTS = {".xlsx", ".xlsm"}
MAX_CELLS_PER_ROW = 40


def load_registry() -> dict:
    if not REGISTRY_PATH.exists():
        return {"active": "", "bases": []}
    return json.loads(REGISTRY_PATH.read_text(encoding="utf-8-sig"))


def resolve_base(raw: str | None) -> Path:
    registry = load_registry()
    if not raw:
        raw = registry.get("active", "")
    for entry in registry.get("bases", []):
        if entry.get("id") == raw:
            path = Path(str(entry.get("path", "")))
            return path if path.is_absolute() else (ROOT / path).resolve()
    candidate = Path(raw).expanduser()
    if candidate.exists():
        return candidate.resolve()
    raise SystemExit(f"[失败] 找不到底座：{raw or '(未设置 active)'}")


def iter_files(base: Path, exts: set[str]):
    for path in base.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in exts:
            continue
        if any(part in SKIP_DIRS for part in path.parts):
            continue
        yield path


def iter_xlsx_rows(path: Path):
    """逐行产出 (sheet, row_no, 行文本)。没有 openpyxl 或读不了就静默跳过。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return
    try:
        for sheet in workbook.worksheets:
            for row_no, row in enumerate(
                sheet.iter_rows(max_col=MAX_CELLS_PER_ROW, values_only=True), start=1
            ):
                cells = [str(cell).strip() for cell in row if cell not in (None, "")]
                if cells:
                    yield sheet.title, row_no, " | ".join(cells)
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="内容底座全文检索")
    parser.add_argument("terms", nargs="*", help="关键词")
    parser.add_argument("--base", default="", help="底座 id 或路径；默认 active")
    parser.add_argument("--ext", action="append", default=[], help="扩展名，可重复；默认常见文本类型")
    parser.add_argument("--limit", type=int, default=80, help="最多输出多少条命中")
    parser.add_argument("--context", type=int, default=0, help="命中行前后各显示多少行")
    parser.add_argument("--any", action="store_true", help="任一关键词命中即算")
    parser.add_argument("--case-sensitive", action="store_true")
    parser.add_argument("--files", action="store_true", help="只列出文件，不搜索")
    parser.add_argument("--no-excel", action="store_true", help="跳过 Excel 单元格检索")
    args = parser.parse_args()

    base = resolve_base(args.base or None)
    exts = {e if e.startswith(".") else "." + e for e in args.ext} if args.ext else DEFAULT_EXTS
    files = list(iter_files(base, exts))
    excel_files = [] if args.no_excel else list(iter_files(base, EXCEL_EXTS))

    if args.files:
        for path in files + excel_files:
            print(path.relative_to(base).as_posix())
        print(f"\n共 {len(files) + len(excel_files)} 个文件（文本 {len(files)} / 表格 {len(excel_files)}）")
        return

    if not args.terms:
        parser.error("请提供至少一个关键词，或使用 --files")

    terms = args.terms if args.case_sensitive else [t.lower() for t in args.terms]
    hits = 0
    scanned = 0
    reached = False

    def is_match(text: str) -> bool:
        haystack = text if args.case_sensitive else text.lower()
        return any(t in haystack for t in terms) if args.any else all(t in haystack for t in terms)

    for path in files:
        scanned += 1
        try:
            lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
        except OSError:
            continue
        rel = path.relative_to(base).as_posix()
        for idx, line in enumerate(lines):
            if not is_match(line):
                continue
            hits += 1
            print(f"{rel}:{idx + 1}: {line.strip()[:240]}")
            if args.context:
                start = max(0, idx - args.context)
                end = min(len(lines), idx + args.context + 1)
                for j in range(start, end):
                    if j == idx:
                        continue
                    print(f"    {rel}:{j + 1}: {lines[j].strip()[:200]}")
            if hits >= args.limit:
                reached = True
                break
        if reached:
            break

    if not reached:
        for path in excel_files:
            scanned += 1
            rel = path.relative_to(base).as_posix()
            for sheet, row_no, text in iter_xlsx_rows(path):
                if not is_match(text):
                    continue
                hits += 1
                print(f'{rel}::{sheet}!A{row_no}: {text[:240]}')
                if hits >= args.limit:
                    reached = True
                    break
            if reached:
                break

    if reached:
        print(f"\n[截断] 已达到 --limit {args.limit}")
    print(f"扫描 {scanned} 个文件（含 {len(excel_files)} 个表格），命中 {hits} 条")


if __name__ == "__main__":
    main()
