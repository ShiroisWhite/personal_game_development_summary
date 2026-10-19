#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel 配置表查看工具（只读）。

用法：
    python tools/excel_inspect.py sheets <file.xlsx>
    python tools/excel_inspect.py schema <file.xlsx> [--sheet NAME] [--rows 8]
    python tools/excel_inspect.py sample <file.xlsx> --sheet NAME [--rows 10]
    python tools/excel_inspect.py search <file.xlsx> --keyword 建造
    python tools/excel_inspect.py find <dir> --keyword 建造 [--limit 50]

说明：
- 支持 .xlsx / .xlsm；.xls 需要先另存为新格式。
- 表头行号因项目而异，本工具只展示原始行，不做假设。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

try:
    from openpyxl import load_workbook
except ImportError:
    print("[失败] 缺少 openpyxl。安装：python -m pip install openpyxl")
    raise SystemExit(1)

MAX_COLS = 40
MAX_CELL = 60


def open_book(path: Path, read_only: bool = True):
    if path.suffix.lower() == ".xls":
        raise SystemExit("[失败] 不支持 .xls，请先另存为 .xlsx")
    if not path.exists():
        raise SystemExit(f"[失败] 文件不存在：{path}")
    return load_workbook(filename=str(path), read_only=read_only, data_only=True)


def cell_text(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", "\\n").replace("|", "\\|")
    if len(text) > MAX_CELL:
        text = text[:MAX_CELL] + "…"
    return text


def print_rows(ws, limit: int, start: int = 1) -> None:
    for row_idx, row in enumerate(ws.iter_rows(min_row=start, max_row=start + limit - 1), start=start):
        values = [cell_text(c.value) for c in row[:MAX_COLS]]
        if not any(values):
            continue
        print(f"{row_idx:>4} | " + " | ".join(values))


def cmd_sheets(args) -> None:
    wb = open_book(Path(args.file))
    print(f"文件：{Path(args.file).resolve()}")
    print(f"Sheet 数量：{len(wb.sheetnames)}\n")
    for name in wb.sheetnames:
        ws = wb[name]
        dims = ""
        try:
            dims = f"{ws.max_row} 行 × {ws.max_column} 列"
        except Exception:
            dims = "（只读模式无法获取尺寸）"
        print(f"- {name}  ({dims})")


def cmd_schema(args) -> None:
    wb = open_book(Path(args.file))
    names = [args.sheet] if args.sheet else wb.sheetnames
    for name in names:
        if name not in wb.sheetnames:
            print(f"[跳过] 不存在 Sheet：{name}")
            continue
        ws = wb[name]
        print(f"\n## {name}")
        try:
            print(f"尺寸：{ws.max_row} 行 × {ws.max_column} 列")
        except Exception:
            pass
        print("\n前几行原始内容（行号 | 值）：")
        print_rows(ws, args.rows)
        print()


def cmd_sample(args) -> None:
    wb = open_book(Path(args.file))
    if args.sheet not in wb.sheetnames:
        raise SystemExit(f"[失败] 不存在 Sheet：{args.sheet}")
    ws = wb[args.sheet]
    print(f"Sheet：{args.sheet}")
    print_rows(ws, args.rows, start=args.start)


def cmd_search(args) -> None:
    wb = open_book(Path(args.file))
    keyword = args.keyword
    if not args.case_sensitive:
        keyword = keyword.lower()
    hits = 0
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                text = str(value)
                haystack = text if args.case_sensitive else text.lower()
                if keyword in haystack:
                    hits += 1
                    print(f"{name}!{cell.coordinate}: {cell_text(value)}")
                    if hits >= args.limit:
                        print(f"\n[截断] 已达到 --limit {args.limit}")
                        return
    print(f"\n命中 {hits} 条")


def cmd_find(args) -> None:
    root = Path(args.dir)
    if not root.exists():
        raise SystemExit(f"[失败] 目录不存在：{root}")
    keyword = args.keyword if args.case_sensitive else args.keyword.lower()
    files = []
    for path in root.rglob("*"):
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            files.append(path)
    print(f"扫描 {len(files)} 个 Excel 文件\n")
    matched = 0
    for path in files:
        try:
            wb = load_workbook(filename=str(path), read_only=True, data_only=True)
        except Exception as exc:
            continue
        hit_sheets = []
        for name in wb.sheetnames:
            haystack = name if args.case_sensitive else name.lower()
            if keyword in haystack:
                hit_sheets.append(name + "（Sheet 名）")
                continue
            ws = wb[name]
            for row in ws.iter_rows(min_row=1, max_row=args.head_rows):
                found = False
                for cell in row:
                    if cell.value is None:
                        continue
                    text = str(cell.value)
                    if keyword in (text if args.case_sensitive else text.lower()):
                        hit_sheets.append(f"{name}!{cell.coordinate}")
                        found = True
                        break
                if found:
                    break
        if hit_sheets:
            matched += 1
            print(f"- {path}")
            for item in hit_sheets[:5]:
                print(f"    {item}")
            if matched >= args.limit:
                print(f"\n[截断] 已达到 --limit {args.limit}")
                return
    print(f"\n命中 {matched} 个文件")


def main() -> None:
    parser = argparse.ArgumentParser(description="Excel 配置表查看工具（只读）")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("sheets", help="列出 Sheet")
    p.add_argument("file")
    p.set_defaults(func=cmd_sheets)

    p = sub.add_parser("schema", help="查看表头与前几行")
    p.add_argument("file")
    p.add_argument("--sheet", default="")
    p.add_argument("--rows", type=int, default=6)
    p.set_defaults(func=cmd_schema)

    p = sub.add_parser("sample", help="查看指定 Sheet 的样例行")
    p.add_argument("file")
    p.add_argument("--sheet", required=True)
    p.add_argument("--rows", type=int, default=10)
    p.add_argument("--start", type=int, default=1)
    p.set_defaults(func=cmd_sample)

    p = sub.add_parser("search", help="在单个文件中搜索")
    p.add_argument("file")
    p.add_argument("--keyword", required=True)
    p.add_argument("--limit", type=int, default=50)
    p.add_argument("--case-sensitive", action="store_true")
    p.set_defaults(func=cmd_search)

    p = sub.add_parser("find", help="在目录中查找包含关键词的文件")
    p.add_argument("dir")
    p.add_argument("--keyword", required=True)
    p.add_argument("--limit", type=int, default=50)
    p.add_argument("--head-rows", type=int, default=5)
    p.add_argument("--case-sensitive", action="store_true")
    p.set_defaults(func=cmd_find)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
